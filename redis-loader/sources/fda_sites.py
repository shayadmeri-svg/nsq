"""FDA signals about Indian manufacturing sites — three public sources:

* fda_establishments  Drug Establishments Current Registration Site (DECRS)
                      bulk file: every establishment registered with FDA,
                      filtered to India (FEI, DUNS, address, business operations).
* fda_import_alerts   Import Alert 66-40 (drug GMP) red list — Indian firms whose
                      products are subject to detention without examination.
* fda_recalls         openFDA drug enforcement reports for firms in India.

The DECRS column names are not published, so the parser finds columns by
name (firm, FEI, DUNS, address, city, state, zip, country, operations,
expiration) and logs the header it saw.
"""

from __future__ import annotations

import csv
import html
import io
import json
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Optional

from .common import Ctx, Unreachable, download, http_get, write_normalized

DECRS = {
    "title": "FDA establishment registrations (DECRS)",
    "publisher": "U.S. Food and Drug Administration",
    "url": "https://www.accessdata.fda.gov/cder/drls_reg.zip",
    "page": "https://www.fda.gov/drugs/drug-approvals-and-databases/drug-establishments-current-registration-site",
    "cadence": "daily file; refreshed weekly here",
    "feeds": ["FDA-registered Indian sites", "FEI / DUNS", "Business operations"],
}
IMPORT_ALERT = {
    "title": "FDA Import Alert 66-40",
    "publisher": "U.S. Food and Drug Administration",
    "url": "https://www.accessdata.fda.gov/cms_ia/importalert_189.html",
    "page": "https://www.accessdata.fda.gov/cms_ia/importalert_189.html",
    "cadence": "daily",
    "feeds": ["Indian firms on the drug-GMP red list"],
}
RECALLS = {
    "title": "openFDA drug recalls (India)",
    "publisher": "U.S. Food and Drug Administration (openFDA)",
    "url": "https://api.fda.gov/drug/enforcement.json",
    "page": "https://open.fda.gov/apis/drug/enforcement/",
    "cadence": "weekly",
    "feeds": ["US recalls of Indian-made drugs", "Recall reasons / class"],
}


def norm_company(name: str) -> str:
    """Loose company key for cross-source matching."""
    n = (name or "").lower()
    n = re.sub(r"\bm/s\.?\s*", " ", n)
    n = re.sub(r"[^a-z0-9 ]+", " ", n)
    n = re.sub(r"\b(private|pvt|limited|ltd|llp|inc|co|company|corporation|corp|the|india|industries|unit|plant|lab|labs|laboratories|laboratory)\b", " ", n)
    return re.sub(r"\s+", " ", n).strip()


# --- DECRS ------------------------------------------------------------------------

_COLS = {
    "name": ("firm_name", "establishment_name", "firm", "name", "registrant_name"),
    "fei": ("fei_number", "fei", "fei_no"),
    "duns": ("duns_number", "duns", "duns_no"),
    "address": ("address_line_1", "address", "street", "address1", "firm_address"),
    "address2": ("address_line_2", "address2"),
    "city": ("city",),
    "state": ("state", "state_province", "province"),
    "postal": ("zip_code", "postal_code", "zip", "postal"),
    "country": ("country_code", "country", "country_name", "iso_country_code"),
    "operations": ("business_operations", "operations", "business_operation", "establishment_operations"),
    "expiration": ("expiration_date", "registration_expiration_date", "expiration", "reg_expiry_date"),
    "registrant": ("registrant_name", "registrant"),
}


def _norm_h(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", (h or "").strip().lower()).strip("_")


def _pick(header: list[str]) -> dict[str, int]:
    out: dict[str, int] = {}
    for field, cands in _COLS.items():
        for c in cands:
            if c in header:
                out[field] = header.index(c)
                break
        else:
            for i, h in enumerate(header):
                if any(h.startswith(c) for c in cands):
                    out[field] = i
                    break
    return out


def _table_rows(path: Path) -> tuple[list[str], list[list[str]]]:
    members: list[tuple[str, bytes]] = []
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as zf:
            members = [(n, zf.read(n)) for n in zf.namelist() if not n.endswith("/")]
    else:
        members = [(path.name, path.read_bytes())]
    name, blob = max(members, key=lambda m: len(m[1]))
    low = name.lower()
    if low.endswith(".xlsx"):
        from openpyxl import load_workbook  # optional dependency
        ws = load_workbook(io.BytesIO(blob), read_only=True).active
        rows = [[("" if v is None else str(v)).strip() for v in row] for row in ws.iter_rows(values_only=True)]
    elif low.endswith(".xls"):
        import pandas as pd  # needs xlrd for .xls
        df = pd.read_excel(io.BytesIO(blob), dtype=str).fillna("")
        rows = [list(df.columns)] + df.values.tolist()
    else:
        text = blob.decode("utf-8", errors="replace") if blob[:3] != b"\xef\xbb\xbf" else blob[3:].decode("utf-8", errors="replace")
        sample = text[:20000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters="|\t,~;")
            delim = dialect.delimiter
        except csv.Error:
            delim = "|" if sample.count("|") > sample.count("\t") else "\t"
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    rows = [r for r in rows if any(r)]
    return [_norm_h(h) for h in rows[0]], rows[1:]


def parse_decrs(path: Path, log=print) -> dict[str, dict[str, Any]]:
    header, rows = _table_rows(path)
    cols = _pick(header)
    log(f"DECRS header: {header[:20]}")
    if "name" not in cols:
        raise RuntimeError(f"DECRS layout not recognised (no firm-name column in {header[:20]})")
    out: dict[str, dict[str, Any]] = {}

    def g(r: list[str], f: str) -> str:
        i = cols.get(f)
        return r[i].strip() if i is not None and i < len(r) else ""

    for r in rows:
        country = g(r, "country").upper()
        joined = " ".join(r).upper()
        if not (country in {"IN", "IND", "INDIA"} or (not country and " INDIA" in joined)):
            continue
        fei = g(r, "fei") or f"noFEI-{len(out)}"
        e = out.setdefault(fei, {
            "fei": g(r, "fei"), "duns": g(r, "duns"), "name": g(r, "name"),
            "address": ", ".join(x for x in (g(r, "address"), g(r, "address2")) if x),
            "city": g(r, "city"), "state": g(r, "state"), "postal": g(r, "postal"),
            "operations": set(), "expiration": g(r, "expiration"), "registrant": g(r, "registrant"),
        })
        for op in re.split(r"\s*[;,]\s*", g(r, "operations")):
            if op:
                e["operations"].add(op.strip().lower())
    for e in out.values():
        e["operations"] = sorted(e["operations"])
        e["company_key"] = norm_company(e["name"])
    return out


def run_establishments(ctx: Ctx) -> int:
    path = ctx.from_file or download(ctx, DECRS["url"], "drls_reg.zip")
    data = parse_decrs(path, ctx.log)
    if not data:
        raise RuntimeError("No Indian establishments found in the DECRS file — check the logged header.")
    write_normalized(ctx, DECRS, data, len(data))
    return len(data)


# --- Import alert 66-40 -----------------------------------------------------------

class _Text(HTMLParser):
    """HTML -> lines, with headings marked '##' and bold marked '**'."""

    def __init__(self) -> None:
        super().__init__()
        self.lines: list[str] = []
        self.cur: list[str] = []
        self.mark = ""

    def _flush(self) -> None:
        t = re.sub(r"\s+", " ", "".join(self.cur)).strip()
        if t:
            self.lines.append(self.mark + t)
        self.cur, self.mark = [], ""

    def handle_starttag(self, tag, attrs):
        if tag in ("h1", "h2", "h3", "h4", "h5", "p", "div", "br", "li", "tr", "table"):
            self._flush()
        if tag in ("h2", "h3", "h4", "h5"):
            self.mark = "## "
        elif tag in ("b", "strong") and not self.cur:
            self.mark = self.mark or "** "

    def handle_endtag(self, tag):
        if tag in ("h1", "h2", "h3", "h4", "h5", "p", "div", "li", "tr", "b", "strong"):
            if tag in ("b", "strong") and self.mark == "** ":
                self._flush()
            elif tag not in ("b", "strong"):
                self._flush()

    def handle_data(self, data):
        self.cur.append(data)


_DATE = re.compile(r"(\d{1,2}/\d{1,2}/\d{4})")
_FEI = re.compile(r"FEI\s*(?:#|No\.?|Number)?\s*:?\s*(\d{6,12})", re.I)


def parse_import_alert(text_html: str, country: str = "INDIA") -> dict[str, dict[str, Any]]:
    p = _Text()
    p.feed(html.unescape(text_html))
    p._flush()
    lines = p.lines
    out: dict[str, dict[str, Any]] = {}
    in_country = False
    cur: Optional[dict[str, Any]] = None
    for ln in lines:
        if ln.startswith("## "):
            head = ln[3:].strip().upper()
            if head == country:
                in_country = True
                continue
            if in_country and re.fullmatch(r"[A-Z ,.'()&-]{3,40}", head) and head != country:
                break  # next country
            continue
        if not in_country:
            continue
        if ln.startswith("** "):
            name = ln[3:].strip()
            if _DATE.fullmatch(name) or name.lower().startswith(("date published", "product", "desc", "notes", "problem")):
                if cur is not None and _DATE.search(name):
                    cur["dates"].append(_DATE.search(name).group(1))
                continue
            cur = out.setdefault(norm_company(name) + "|" + str(len(out)), {"name": name, "address": [], "fei": "", "dates": [], "notes": []})
            continue
        if cur is None:
            continue
        m = _FEI.search(ln)
        if m:
            cur["fei"] = m.group(1)
        d = _DATE.search(ln)
        if d and "publish" in ln.lower():
            cur["dates"].append(d.group(1))
        elif len(cur["address"]) < 4 and not ln.lower().startswith(("date", "product", "desc", "notes", "problem", "*")) and not d:
            cur["address"].append(ln)
        elif len(cur["notes"]) < 3:
            cur["notes"].append(ln[:240])
    res = {}
    for e in out.values():
        key = e["fei"] or norm_company(e["name"])
        res[key] = {"name": e["name"], "fei": e["fei"], "address": ", ".join(e["address"]),
                    "company_key": norm_company(e["name"]), "dates": sorted(set(e["dates"])), "notes": e["notes"]}
    return res


def run_import_alerts(ctx: Ctx) -> int:
    if ctx.from_file:
        text = ctx.from_file.read_text(encoding="utf-8", errors="replace")
    else:
        path = download(ctx, IMPORT_ALERT["url"], "importalert_66-40.html")
        text = path.read_text(encoding="utf-8", errors="replace")
    data = parse_import_alert(text)
    if not data:
        raise RuntimeError("No Indian firms parsed from Import Alert 66-40 — the page layout may have changed.")
    write_normalized(ctx, IMPORT_ALERT, data, len(data))
    return len(data)


# --- openFDA recalls -------------------------------------------------------------

def parse_recalls(results: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for r in results:
        firm = r.get("recalling_firm", "")
        key = norm_company(firm)
        if not key:
            continue
        e = out.setdefault(key, {"name": firm, "company_key": key, "city": r.get("city", ""), "state": r.get("state", ""),
                                 "recalls": 0, "class_i": 0, "last": "", "items": []})
        e["recalls"] += 1
        if (r.get("classification") or "").strip() == "Class I":
            e["class_i"] += 1
        d = r.get("recall_initiation_date") or r.get("report_date") or ""
        if d > e["last"]:
            e["last"] = d
        if len(e["items"]) < 8:
            e["items"].append({
                "date": d, "classification": r.get("classification", ""), "status": r.get("status", ""),
                "product": (r.get("product_description") or "")[:200], "reason": (r.get("reason_for_recall") or "")[:300],
                "recall_number": r.get("recall_number", ""),
            })
    for e in out.values():
        e["items"].sort(key=lambda x: x["date"], reverse=True)
    return out


def run_recalls(ctx: Ctx) -> int:
    results: list[dict[str, Any]] = []
    if ctx.from_file:
        payload = json.loads(ctx.from_file.read_text(encoding="utf-8"))
        results = payload.get("results", payload if isinstance(payload, list) else [])
    else:
        skip, page = 0, 1000
        while skip < 25000:
            raw = http_get(RECALLS["url"], params={"search": 'country:"India"', "limit": str(page), "skip": str(skip)},
                           accept="application/json", timeout=120)
            body = json.loads(raw.decode("utf-8"))
            batch = body.get("results") or []
            results.extend(batch)
            total = (body.get("meta") or {}).get("results", {}).get("total", 0)
            skip += page
            if not batch or skip >= total:
                break
        ctx.log(f"openFDA returned {len(results):,} enforcement reports")
    data = parse_recalls(results)
    write_normalized(ctx, RECALLS, data, len(data), extra={"reports": len(results)})
    return len(data)
